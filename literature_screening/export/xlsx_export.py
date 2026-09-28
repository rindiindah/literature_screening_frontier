"""
export/xlsx_export.py
========================
Ekspor dataset screening ke XLSX (Bagian 16) dengan data validation
dropdown pada kolom `decision` (Include/Exclude) dan `exclusion_reason`
agar peneliti mudah screening langsung di Excel (Option B, Bagian 20).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import settings
from export.csv_export import OUTPUT_COLUMNS


def export_xlsx(records: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Screening"

    ws.append(OUTPUT_COLUMNS)

    for rec in records:
        row = []
        for col in OUTPUT_COLUMNS:
            value = rec.get(col)
            if isinstance(value, list):
                value = "; ".join(value)
            row.append(value)
        ws.append(row)

    n_rows = len(records) + 1  # +1 header

    decision_col_idx = OUTPUT_COLUMNS.index("decision") + 1
    reason_col_idx = OUTPUT_COLUMNS.index("exclusion_reason") + 1
    decision_col_letter = get_column_letter(decision_col_idx)
    reason_col_letter = get_column_letter(reason_col_idx)

    decision_dv = DataValidation(
        type="list",
        formula1=f'"{",".join(settings.DECISION_VALUES)}"',
        allow_blank=True,
        showDropDown=False,
    )
    decision_dv.error = "Hanya boleh Include atau Exclude (tanpa Maybe)."
    decision_dv.errorTitle = "Nilai tidak valid"
    ws.add_data_validation(decision_dv)
    decision_dv.add(f"{decision_col_letter}2:{decision_col_letter}{n_rows}")

    reason_dv = DataValidation(
        type="list",
        formula1=f'"{",".join(settings.EXCLUSION_REASONS)}"',
        allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(reason_dv)
    reason_dv.add(f"{reason_col_letter}2:{reason_col_letter}{n_rows}")

    # Lebar kolom wajar agar nyaman dibaca
    widths = {
        "title": 45, "abstract": 60, "authors": 30, "journal": 25,
        "url": 25, "notes": 25, "source_database": 18,
    }
    for i, col in enumerate(OUTPUT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 15)

    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[XLSX Export] {len(records)} baris ditulis ke {output_path} (dropdown decision/exclusion_reason aktif)")
